#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""重建「可重建输入」——被 tools\\deploy_er.ps1 的孤儿分支部署清理误删的 .gitignore 登记项。

根因：部署流程 `checkout --orphan gh-pages` 前后清工作区（AGENTS.md §部署「清工作区先移出 ER-source」）
      只保住了 ER-source，未跟踪/被忽略的构建输入被一并删除。
重建依据（全部来自「权威源 + 已提交产物」的反演，不引入新数据）：
  1. 招式表_可学总表.csv                      ← ER-source\\gameDataV2.65beta.json（levelUp/tutor/egg/TMHM）
                                               + 招式表_招式数值.csv / 招式表_宝可梦基础.csv（中文名等展示列）
                                               行序与 build_movesets.py 完全一致（sid 升序：升级→教学→蛋→TMHM）
  2. ER2.65简汉化\\ER2.65beta版图鉴v0.3.xlsm   ← 配招工具_data.js 既有产物
        · Sheet1 克制方阵（ERDATA.types + ERDATA.matchup 原样回填，21×21）
        · 特性   id → 中文名 / 中文描述（ERDATA.abilities；v0.3 侧含 ABI_ALIAS 旧译名，行序复刻别名插入顺序）
  3. ER2.65简汉化\\分类.xlsx                   ← ERDATA.coreNotes 原样回填（Sheet1，行序一致）
  4. ER2.5正式版图鉴v0.5.xlsm                  ← 「特性」表（正式版译名，合并时优先，故直接取最终名）

验收：跑 build_tool_data.py 后，除新增的 usagePrior / usageMeta 外，产物须与既有
      配招工具_data.js **逐字段一致**（见 tools\\check_rebuild_faithful.py）。
"""
import csv, json, os

BASE = r'D:\game\elite-redux'
DGJ = BASE + r'\ER2.65简汉化'
TYPE_ORDER = ['一般', '格斗', '火', '冰', '电', '虫', '飞行', '钢', '草', '地面',
              '毒', '恶', '水', '超能力', '岩石', '龙', '幽灵', '妖精', '星晶', '无', '神秘']
HDR = ['宝可梦编号', '宝可梦中文', '宝可梦英文', '属性1', '属性2',
       '种族HP', '种族攻击', '种族防御', '种族特攻', '种族特防', '种族速度',
       '特性1-3', '天性1-3', '招式编号', '招式中文', '招式英文', '招式属性', '分类',
       '威力', '命中', 'PP', '先制', '获得方式', '等级', '招式描述']

# ---------- 0. 读取既有产物（oracle） ----------
_txt = open(BASE + r'\配招工具_data.js', encoding='utf-8').read()
_i = _txt.index('var ERDATA = ') + len('var ERDATA = ')
D, _ = json.JSONDecoder().raw_decode(_txt[_i:])
print('oracle: 配招工具_data.js 顶层键 %d 个' % len(D))
gd = json.load(open(BASE + r'\ER-source\gameDataV2.65beta.json', encoding='utf-8'))
sp_map = {s['id']: s for s in gd['species']}
gd_abi = {a['id']: a for a in gd['abilities']}

# ---------- 1. 招式表_可学总表.csv ----------
sp_base, mv_num = {}, {}
for r in csv.reader(open(BASE + r'\招式表_宝可梦基础.csv', encoding='utf-8-sig')):
    if r and r[0] != '宝可梦编号':
        sp_base[r[0]] = r
for r in csv.reader(open(BASE + r'\招式表_招式数值.csv', encoding='utf-8-sig')):
    if r and r[0] != '招式编号':
        mv_num[r[0]] = r


def mv_col(mid, idx):
    r = mv_num.get(str(mid))
    return r[idx] if r and len(r) > idx else ''


rows, cnt = [], {'升级': 0, '教学': 0, '蛋': 0, 'TM/HM': 0}


def emit(sid, mid, way, lv=''):
    b = sp_base.get(str(sid), [])
    g = lambda i: b[i] if len(b) > i else ''
    rows.append([sid, g(1), g(2), g(3), g(4), g(5), g(6), g(7), g(8), g(9), g(10), g(11), g(12),
                 mid, mv_col(mid, 1), mv_col(mid, 2), mv_col(mid, 3), mv_col(mid, 4),
                 mv_col(mid, 5), mv_col(mid, 6), mv_col(mid, 7), mv_col(mid, 8),
                 way, lv, mv_col(mid, 10)])
    cnt[way] += 1


for sid in sorted(sp_map):
    if sid < 0:
        continue
    s = sp_map[sid]
    for e in (s.get('levelUpMoves') or []):
        emit(sid, e['id'] if isinstance(e, dict) else e, '升级', (e.get('lv') if isinstance(e, dict) else '') or '')
    for e in (s.get('tutor') or []):
        emit(sid, e['id'] if isinstance(e, dict) else e, '教学')
    for e in (s.get('eggMoves') or []):
        emit(sid, e['id'] if isinstance(e, dict) else e, '蛋')
    for e in (s.get('TMHMMoves') or []):
        emit(sid, e['id'] if isinstance(e, dict) else e, 'TM/HM')

with open(BASE + r'\招式表_可学总表.csv', 'w', encoding='utf-8-sig', newline='') as f:
    w = csv.writer(f)
    w.writerow(HDR)
    w.writerows(rows)
_exp = {'升级': 40299, '教学': 72281, '蛋': 0, 'TM/HM': 0}
print('招式表_可学总表.csv: %d 行 %s (期望 112580 / %s) %s' %
      (len(rows), cnt, _exp, 'OK' if (len(rows) == 112580 and cnt == _exp) else '★不符'))
assert len(rows) == 112580 and cnt == _exp, '可学总表行数与基线不符'

# ---------- 2. ER2.65beta版图鉴v0.3.xlsm（Sheet1 克制 + 特性中文） ----------
import openpyxl
from openpyxl import Workbook

# 2a. 特性表：v0.3 侧旧译名（= ABI_ALIAS 的键，按产物插入顺序复刻 v0.5 行序）
alias = D['abiAlias']                       # {v0.3 旧译名: aid}
v03_name = {aid: nm for nm, aid in alias.items()}
order = list(alias.values()) + [a for a in sorted(gd_abi) if a not in set(alias.values())]
desc_by_aid = {}
for row in D['abilities']:                  # [str(id), en, zh, desc]
    aid = int(row[0])
    en = row[1]
    zh = row[2]
    gdesc = (gd_abi.get(aid) or {}).get('desc', '')
    cdesc = row[3] if row[3] != gdesc else ''      # desc = cdesc or gameData desc 的反演
    desc_by_aid[aid] = (zh, cdesc)


def write_abi_wb(path, name_of):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = '特性'
    # 列布局：0=id 1=(空) 2=英文名 3=中文名 4=(空) 5=中文描述 6=(空) 7=描述续行
    for aid in order:
        zh, cdesc = desc_by_aid.get(aid, ('', ''))
        if not zh:
            continue
        head, tail = (cdesc.split('\n', 1) + [''])[:2] if '\n' in cdesc else (cdesc, '')
        ws.append([aid, '', (gd_abi.get(aid) or {}).get('name', ''), name_of(aid), '', head, '', tail])
    wb.save(path)


write_abi_wb(DGJ + r'\ER2.65beta版图鉴v0.3.xlsm', lambda a: v03_name.get(a) or desc_by_aid.get(a, ('', ''))[0])

# 2b. 同一 workbook 再补 Sheet1 克制方阵（重开写入，保留特性表）
wb = openpyxl.load_workbook(DGJ + r'\ER2.65beta版图鉴v0.3.xlsm')
ws = wb.create_sheet('Sheet1')
ws.append(['', '', *TYPE_ORDER])
ws.append(['', '', *TYPE_ORDER])
for di, dt in enumerate(TYPE_ORDER):
    ws.append(['', dt, *D['matchup'][di]])
wb.save(DGJ + r'\ER2.65beta版图鉴v0.3.xlsm')
print('ER2.65简汉化\\ER2.65beta版图鉴v0.3.xlsm 重建：特性 %d 行 + Sheet1 %dx%d（%d 列头）' %
      (len(order), len(TYPE_ORDER), len(TYPE_ORDER), len(TYPE_ORDER)))

# ---------- 3. ER2.5正式版图鉴v0.5.xlsm（正式版译名，合并优先） ----------
write_abi_wb(BASE + r'\ER2.5正式版图鉴v0.5.xlsm', lambda a: desc_by_aid.get(a, ('', ''))[0])
print('ER2.5正式版图鉴v0.5.xlsm 重建：特性 %d 行' % len(order))

# ---------- 4. ER2.65简汉化\分类.xlsx（coreNotes 原样回填） ----------
wb = Workbook()
ws = wb.active
ws.title = 'Sheet1'
for n in D['coreNotes']:
    ws.append([n['name'], n['cat'], '', n['note'],
               n['atk']['bst'], n['atk']['mul'], n['atk']['pow'], n['atk']['boost'],
               n['def']['bst'], n['def']['mul'], n['def']['boost'], n['def']['resist'],
               n['remark']])
wb.save(DGJ + r'\分类.xlsx')
print('ER2.65简汉化\\分类.xlsx 重建：%d 行' % len(D['coreNotes']))

# ---------- 5. 出处说明 ----------
open(DGJ + r'\重建说明.txt', 'w', encoding='utf-8').write(
    '本目录文件由部署清理误删（.gitignore 登记为可重建输入）。\n'
    '重建工具：tools\\rebuild_source_inputs.py\n'
    '重建来源：ER-source\\gameDataV2.65beta.json + 招式表_招式数值.csv / 招式表_宝可梦基础.csv\n'
    '          + 既有产物 配招工具_data.js（ERDATA.matchup / abilities / coreNotes 原样回填）。\n'
    '重建后验收：build_tool_data.py 产物除 usagePrior/usageMeta 外须与既有配招工具_data.js 逐字段一致。\n')
print('重建完成（4 项输入 + 重建说明.txt）')
